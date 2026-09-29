import streamlit as st
from docx import Document
from datetime import date
import io

st.set_page_config(page_title="LegalEase AI", page_icon="⚖️")
st.title("⚖️ LegalEase - AI Legal Document Generator")

TEMPLATES = {
    "Rental Agreement": "RENTAL AGREEMENT\n\nDate: {date}\n\nLANDLORD: {party1_name}, {party1_address}\nTENANT: {party2_name}, {party2_address}\n\nPROPERTY: {property_details}\nRENT: Rs. {rent_amount}/month\nDURATION: {duration} months from {start_date}\nDEPOSIT: Rs. {deposit_amount}\n\nBoth parties agree to the terms above.\n\nLandlord: ___________   Tenant: ___________",
    "NDA": "NON-DISCLOSURE AGREEMENT\n\nDate: {date}\nBetween {party1_name} and {party2_name}\n\nPurpose: {purpose}\nConfidential Info: {property_details}\nDuration: {duration} years\n\nReceiving party shall not disclose confidential information.\n\n{party1_name}: ___________   {party2_name}: ___________",
    "Employment Letter": "EMPLOYMENT OFFER LETTER\n\nDate: {date}\nTo: {party2_name}, {party2_address}\n\nDear {party2_name},\nWe offer you the position of {purpose} at {party1_name}.\nSalary: Rs. {rent_amount}/month\nStart Date: {start_date}\nLocation: {property_details}\n\nSincerely,\n{party1_name}"
}

doc_type = st.sidebar.selectbox("Select Document", list(TEMPLATES.keys()))

party1_name = st.text_input("Party 1 / Company Name")
party1_address = st.text_input("Party 1 Address")
party2_name = st.text_input("Party 2 Name")
party2_address = st.text_input("Party 2 Address")
property_details = st.text_area("Property / Details")
rent_amount = st.text_input("Rent / Salary Amount")
duration = st.text_input("Duration")
start_date = st.date_input("Start Date", date.today())
deposit_amount = st.text_input("Deposit")
purpose = st.text_input("Purpose / Job Role")

if st.button("Generate Document"):
    doc_text = TEMPLATES[doc_type].format(
        date=str(date.today()), party1_name=party1_name, party1_address=party1_address,
        party2_name=party2_name, party2_address=party2_address, property_details=property_details,
        rent_amount=rent_amount, duration=duration, start_date=start_date,
        deposit_amount=deposit_amount, purpose=purpose
    )
    st.text_area("Document Preview", doc_text, height=400)
    
    document = Document()
    document.add_heading(doc_type, 0)
    document.add_paragraph(doc_text)
    bio = io.BytesIO()
    document.save(bio)
    st.download_button("Download DOCX", bio.getvalue(), file_name=f"{doc_type}.docx")
