from io import BytesIO

from docx import Document
from fpdf import FPDF

from utils.document_utils import sanitize_text


def create_txt(text):

    return sanitize_text(text)


def create_docx(text, document_type):

    document = Document()

    title = document.add_heading(
        document_type.upper(),
        level=0
    )

    document.add_paragraph(
        sanitize_text(text)
    )

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer


def create_pdf(text, document_type):

    pdf = FPDF()

    pdf.add_page()
    pdf.set_auto_page_break(
        auto=True,
        margin=15
    )

    pdf.set_font("Arial", "B", 16)
    pdf.cell(
        0,
        10,
        document_type.upper(),
        ln=True,
        align="C"
    )

    pdf.ln(10)

    pdf.set_font("Arial", size=11)

    clean_text = sanitize_text(text)

    for paragraph in clean_text.split("\n"):

        if paragraph.strip():
            pdf.multi_cell(
                0,
                7,
                paragraph
            )

    buffer = BytesIO()
    pdf.output(buffer)

    buffer.seek(0)

    return buffer
